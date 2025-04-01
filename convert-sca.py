import json
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from openpyxl import load_workbook

# --- Chuẩn hóa dữ liệu từ Snyk ---
def normalize_snyk(data):
    try:
        result = []
        for vuln in data.get("vulnerabilities", []):
            result.append({
                "id": vuln.get("id") or vuln.get("identifiers", {}).get("CVE", [None])[0],
                "package": vuln.get("moduleName") or vuln.get("packageName"),
                "version": vuln.get("version"),
                "severity": vuln.get("severity"),
                "remediation": ", ".join(vuln.get("fixedIn", [])),
                "source": "snyk",
                "reference": "./snyk.html"
            })
        return result
    except Exception:
        return []

# --- Chuẩn hóa dữ liệu từ Trivy ---
def normalize_trivy(data):
    try:
        result = []
        for entry in data.get("Results", []):
            for vuln in entry.get("Vulnerabilities", []):
                result.append({
                    "id": vuln.get("VulnerabilityID"),
                    "package": vuln.get("PkgName"),
                    "version": vuln.get("InstalledVersion"),
                    "severity": vuln.get("Severity"),
                    "remediation": vuln.get("FixedVersion", ""),
                    "source": "trivy",
                    "reference": "./trivy-fs.html"
                })
        return result
    except Exception:
        return []

# --- Loại trùng theo ID ---
def deduplicate_by_id(vulns):
    seen = set()
    deduped = []
    for v in vulns:
        if v["id"] and v["id"] not in seen:
            seen.add(v["id"])
            deduped.append(v)
    return deduped

# --- Xuất Excel ---
def export_to_excel(vulns, excel_out_path, sheet_name="SCA"):
    # Thêm số thứ tự
    for i, item in enumerate(vulns, start=1):
        item["#"] = i

    # Nếu danh sách rỗng → tạo bảng rỗng đúng định dạng
    if not vulns:
        columns = ["#", "ID", "Package", "Version", "Severity", "Remediation", "Source", "Reference"]
        df = pd.DataFrame(columns=columns)
    else:
        # Đảm bảo đúng thứ tự cột
        df = pd.DataFrame(vulns)[["#", "id", "package", "version", "severity", "remediation", "source", "reference"]]
        df.columns = ["#", "ID", "Package", "Version", "Severity", "Remediation", "Source", "Reference"]

    # Ghi Excel
    df.to_excel(excel_out_path, sheet_name=sheet_name, index=False)

    # Format đẹp
    wb = load_workbook(excel_out_path)
    ws = wb[sheet_name]

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max_length + 2, 60)
        ws.column_dimensions[col_letter].width = adjusted_width

    wb.save(excel_out_path)

# --- Tổng hợp và xuất ---
def generate_bug_report_excel(snyk_path, trivy_path, output_excel_path):
    try:
        with open(snyk_path, "r", encoding="utf-8") as f:
            snyk_data = json.load(f)
    except Exception:
        snyk_data = {}

    try:
        with open(trivy_path, "r", encoding="utf-8") as f:
            trivy_data = json.load(f)
    except Exception:
        trivy_data = {}

    merged = normalize_snyk(snyk_data) + normalize_trivy(trivy_data)
    filtered = deduplicate_by_id(merged)
    export_to_excel(filtered, output_excel_path, sheet_name="SCA")

# --- Gọi chạy ---
generate_bug_report_excel(
    "report/result/snyk.json",
    "report/result/trivy-fs.json",
    "unique_vulns_filtered.xlsx"
)
