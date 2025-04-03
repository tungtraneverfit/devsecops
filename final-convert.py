
import json
import pandas as pd
from datetime import datetime
import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


today = datetime.now()
today_str = today.strftime("%d-%m-%y")
output_file = f"devsecops-report-{today_str}.xlsx"


def write_to_excel(df, sheet_name):
    print(output_file)
    file_exists = os.path.exists(output_file)
    mode = "a" if file_exists else "w"
    df["Date"] = pd.to_datetime(today)

    if file_exists:
        writer = pd.ExcelWriter(output_file, engine="openpyxl", mode="a", if_sheet_exists="replace")
    else:
        writer = pd.ExcelWriter(output_file, engine="openpyxl", mode="w")

    with writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)

    print(f"✅ Written to sheet '{sheet_name}' in {output_file}")


def convert_sarif_results(json_path, sheet_name):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    results = data.get("runs", [])[0].get("results", [])
    parsed = []
    for result in results:
        for loc in result.get("locations", []):
            region = loc.get("physicalLocation", {}).get("region", {})
            parsed.append({
                "Rule ID": result.get("ruleId"),
                # "Level": result.get("level"),
                "Message": result.get("message", {}).get("text", ""),
                "File": loc.get("physicalLocation", {}).get("artifactLocation", {}).get("uri", ""),
                "Start Line": region.get("startLine", ""),
                "End Line": region.get("endLine", "")
            })
    if not parsed:
        df = pd.DataFrame(columns=[
            "Rule ID", "Message", "File", "Start Line", "End Line",
        ])
        df = df.drop_duplicates()
        df = df.sort_values(by="Rule ID", ascending=True)
    else:
        df = pd.DataFrame(parsed)
    write_to_excel(df, sheet_name)


# def convert_snyk_sast_sarif(json_path, sheet_name="snyk-sast"):
#     with open(json_path, "r", encoding="utf-8") as f:
#         data = json.load(f)

#     rules = {r.get("id"): r for r in data.get("runs", [])[0].get("tool", {}).get("driver", {}).get("rules", [])}
#     results = data.get("runs", [])[0].get("results", [])

#     rows = []
#     for result in results:
#         rule_id = result.get("ruleId", "")
#         rule = rules.get(rule_id, {})
#         message = result.get("message", {}).get("text", "")
#         level = result.get("level", "")

#         for loc in result.get("locations", []):
#             physical = loc.get("physicalLocation", {})
#             artifact = physical.get("artifactLocation", {}).get("uri", "")
#             region = physical.get("region", {})
#             start_line = region.get("startLine", "")
#             end_line = region.get("endLine", "")

#             rows.append({
#                 "Rule ID": rule_id,
#                 "Rule Name": rule.get("name", ""),
#                 "Short Description": rule.get("shortDescription", {}).get("text", ""),
#                 "Full Description": rule.get("fullDescription", {}).get("text", ""),
#                 "Help": rule.get("help", {}).get("text", ""),
#                 "Message": message,
#                 "Severity": level,
#                 "File Path": artifact,
#                 "Line Start": start_line,
#                 "Line End": end_line
#             })

#     if not rows:
#         df = pd.DataFrame(columns=[
#             "Rule ID", "Rule Name", "Short Description", "Full Description",
#             "Help", "Message", "Severity", "File Path", "Line Start", "Line End"
#         ])
#     else:
#         df = pd.DataFrame(rows)

#     write_to_excel(df, sheet_name)


def convert_snyk_vuln_json(json_path, sheet_name):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    vulns = data.get("vulnerabilities", [])
    rows = []
    for v in vulns:
        rows.append({
            "ID": v.get("id"),
            "Title": v.get("title"),
            "Severity": v.get("severity"),
            "Package": v.get("name") or v.get("packageName"),
            "Version": v.get("version"),
            "CVSS": v.get("CVSSv3"),
            "CVEs": ", ".join(v.get("identifiers", {}).get("CVE", [])),
            "Description": v.get("description", "").strip().split("\n")[0],
            "Fixed In": ", ".join(v.get("fixedIn", [])),
            "References": ", ".join([r.get("url") for r in v.get("references", [])])
        })
    if not rows:
        df = pd.DataFrame(columns=[
            "ID", "Title", "Severity", "Package", "Version",
            "CVSS", "CVEs", "Description", "Fixed In", "References"
        ])
    else:
        df = pd.DataFrame(rows)
        df = df.drop_duplicates()
        df = df.sort_values(by="Severity", ascending=True)
    write_to_excel(df, sheet_name)


def convert_trivy_vuln(json_path, sheet_name):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    results = data.get("Results", [])
    rows = []
    for r in results:
        vulns = r.get("Vulnerabilities", [])
        for v in vulns:
            rows.append({
                "VulnerabilityID": v.get("VulnerabilityID"),
                "Target": r.get("Target"),
                "PkgName": v.get("PkgName"),
                "InstalledVersion": v.get("InstalledVersion"),
                "Severity": v.get("Severity"),
                "Title": v.get("Title"),
                "Description": v.get("Description", "").strip().split("\n")[0],
                "FixedVersion": v.get("FixedVersion"),
                "CVSS Score": v.get("CVSS", {}).get("nvd", {}).get("V3Score", ""),
                "CVSS Vector": v.get("CVSS", {}).get("nvd", {}).get("V3Vector", ""),
                "References": ", ".join(v.get("References", []))
            })
    if not rows:
        df = pd.DataFrame(columns=[
            "VulnerabilityID", "Target", "PkgName", "InstalledVersion", "Severity",
            "Description", "FixedVersion", "CVSS Score", "CVSS Vector", "References"
        ])
    else:
        df = pd.DataFrame(rows)
        df = df.drop_duplicates()
        df = df.sort_values(by="Severity", ascending=True)
    write_to_excel(df, sheet_name)


def convert_trivy_k8s(json_path, sheet_name):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    rows = []
    for res in data.get("Resources", []):
        for result in res.get("Results", []):
            for m in result.get("Misconfigurations", []):
                rows.append({
                    "ID": m.get("ID"),
                    "Kind": res.get("Kind"),
                    "Name": res.get("Name"),
                    "Target": result.get("Target"),
                    "Title": m.get("Title"),
                    "Severity": m.get("Severity"),
                    "Message": m.get("Message"),
                    "Resolution": m.get("Resolution"),
                    "References": ", ".join(m.get("References", [])),
                })
    if not rows:
        df = pd.DataFrame(columns=[
            "ID", "Kind", "Name", "Target", "Title",
            "Severity", "Message", "Resolution", "References"
        ])
    else:
        df = pd.DataFrame(rows)
        df = df.drop_duplicates()
        df = df.sort_values(by="Severity", ascending=True)
        
    write_to_excel(df, sheet_name)


def merge_semgrep_and_snyk_to_sast():
    xls = pd.ExcelFile(output_file)

    if "semgrep" not in xls.sheet_names or "snyk-sast" not in xls.sheet_names:
        print("⚠️ 'semgrep' or 'snyk-sast' sheet not found. Skipping SAST merge.")
        return

    df_semgrep = xls.parse("semgrep")
    df_snyk = xls.parse("snyk-sast")

    df_semgrep["source"] = "semgrep"
    df_snyk["source"] = "snyk"

    # sync column names
    df_semgrep.rename(columns={"Rule ID": "rule_id"}, inplace=True)
    df_snyk.rename(columns={"Rule ID": "rule_id"}, inplace=True)

    # Add missing columns
    for col in ["rule_id", "Message", "File", "Start Line", "End Line", "Date"]:
        if col not in df_semgrep.columns:
            df_semgrep[col] = ""
        if col not in df_snyk.columns:
            df_snyk[col] = ""

    selected_cols = ["source", "rule_id", "Message", "File", "Start Line", "End Line", "Date"]

    df_combined = pd.concat([
        df_semgrep[selected_cols],
        df_snyk[selected_cols]
    ], ignore_index=True)

    # sync column names
    df_combined.rename(columns={"rule_id": "Rule ID"}, inplace=True)
    df_combined.rename(columns={"source": "Source"}, inplace=True)

    with pd.ExcelWriter(output_file, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        df_combined.to_excel(writer, sheet_name="SAST", index=False)

    print("✅ Merged 'semgrep' and 'snyk-sast' into sheet 'SAST'")


def merge_trivy_and_snyk_to_sca():
    xls = pd.ExcelFile(output_file)

    if "trivy-fs" not in xls.sheet_names or "snyk" not in xls.sheet_names:
        print("⚠️ 'trivy-fs' or 'snyk' sheet not found. Skipping SCA merge.")
        return

    df_trivy = xls.parse("trivy-fs")
    df_snyk = xls.parse("snyk")

    df_trivy_sca = pd.DataFrame({
        "ID": df_trivy["VulnerabilityID"],
        "Package": df_trivy["PkgName"],
        "Version": df_trivy["InstalledVersion"],
        "Severity": df_trivy["Severity"],
        "Remediation": df_trivy["FixedVersion"],
        "Reference": df_trivy["References"],
        "Date": df_trivy["Date"],
        "Source": "trivy-fs"
    })

    df_snyk_sca = pd.DataFrame({
        "ID": df_snyk["ID"],
        "Package": df_snyk["Package"],
        "Version": df_snyk["Version"],
        "Severity": df_snyk["Severity"],
        "Remediation": df_snyk["Fixed In"],
        "Reference": df_snyk["References"],
        "Date": df_snyk["Date"],
        "Source": "snyk-fs"
    })

    df_sca = pd.concat([df_trivy_sca, df_snyk_sca], ignore_index=True)
    df_sca = df_sca[["Source", "ID", "Package", "Version", "Severity", "Remediation", "Reference","Date"]]

    df_sca["Severity"] = df_sca["Severity"].astype(str).str.upper()
    
    severity_order = {"CRITICAL": 1, "HIGH": 2, "MEDIUM": 3, "LOW": 4, "UNKNOWN": 5}
    df_sca["Severity_Sort"] = df_sca["Severity"].map(severity_order)
    df_sca = df_sca.sort_values(by="Severity_Sort").drop(columns=["Severity_Sort"])

    with pd.ExcelWriter(output_file, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        df_sca.to_excel(writer, sheet_name="SCA", index=False)

    print("✅ Merged 'trivy-fs' and 'snyk-fs' into sheet 'SCA'")


def merge_trivy_and_snyk_to_image_scan():
    xls = pd.ExcelFile(output_file)

    if "trivy-image" not in xls.sheet_names or "snyk-image" not in xls.sheet_names:
        print("⚠️ 'trivy-image' or 'snyk-image' sheet not found. Skipping IMAGE_SCAN merge.")
        return

    df_trivy = xls.parse("trivy-image")
    df_snyk = xls.parse("snyk-image")

    df_trivy_image = pd.DataFrame({
        "Source": "trivy-image",
        "CVE": df_trivy["VulnerabilityID"],
        "Title": df_trivy["Title"],
        "Severity": df_trivy["Severity"],
        "Package": df_trivy["PkgName"],
        "Version": df_trivy["InstalledVersion"],
        "Fixed_In": df_trivy["FixedVersion"],
        "References": df_trivy["References"],
        "Date": df_trivy["Date"],
    })

    df_snyk_image = pd.DataFrame({
        "Source": "snyk-image",
        "CVE": df_snyk["CVEs"],
        "Title": df_snyk["Title"],
        "Severity": df_snyk["Severity"],
        "Package": df_snyk["Package"],
        "Version": df_snyk["Version"],
        "Fixed_In": df_snyk["Fixed In"],
        "References": df_snyk["References"],
        "Date": df_snyk["Date"],
    })

    df_image = pd.concat([df_trivy_image, df_snyk_image], ignore_index=True)
    df_image = df_image[["Source", "CVE", "Title", "Severity", "Package", "Version", "Fixed_In", "References", "Date"]]

    df_image["Severity"] = df_image["Severity"].astype(str).str.upper()
    
    severity_order = {"CRITICAL": 1, "HIGH": 2, "MEDIUM": 3, "LOW": 4, "UNKNOWN": 5}
    df_image["Severity_Sort"] = df_image["Severity"].map(severity_order)
    df_image = df_image.sort_values(by="Severity_Sort").drop(columns=["Severity_Sort"])
    
    with pd.ExcelWriter(output_file, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        df_image.to_excel(writer, sheet_name="IMAGE_SCAN", index=False)

    print("✅ Merged 'trivy-image' and 'snyk-image' into sheet 'IMAGE_SCAN'")


def merge_devsecops_reports(file_today_path, file_yesterday_path, output_path):
    xls_today = pd.read_excel(file_today_path, sheet_name=None)
    xls_yesterday = pd.read_excel(file_yesterday_path, sheet_name=None)

    sheet_names = list(xls_today.keys())
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        for sheet in sheet_names:
            df_today = xls_today[sheet].copy()
            df_yesterday = xls_yesterday[sheet].copy()

            if 'Date' not in df_today.columns or 'Date' not in df_yesterday.columns:
                df_today.to_excel(writer, sheet_name=sheet, index=False)
                continue

            df_today['Date'] = pd.to_datetime(df_today['Date'], errors='coerce')
            df_yesterday['Date'] = pd.to_datetime(df_yesterday['Date'], errors='coerce')

            compare_cols = [col for col in df_today.columns if col != 'Date']

            merged = pd.concat([df_today, df_yesterday], ignore_index=True)
            if 'Severity' in merged.columns:
                severity_order = {"Critical": 1, "High": 2, "Medium": 3, "Low": 4, "Unknown": 5}
                merged["Severity_Sort"] = merged["Severity"].map(severity_order).fillna(999)
                merged.sort_values(by=["Date", "Severity_Sort"], ascending=[False, True], inplace=True)
                merged.drop(columns=["Severity_Sort"], inplace=True)
            else:
                merged.sort_values(by="Date", ascending=False, inplace=True)
            merged = merged.drop_duplicates(subset=compare_cols, keep='first')

            merged.to_excel(writer, sheet_name=sheet, index=False)

    wb = load_workbook(output_path)
    for sheet in sheet_names:
        ws = wb[sheet]
        df_today = xls_today[sheet].copy()
        df_yesterday = xls_yesterday[sheet].copy()

        if 'Date' not in df_today.columns or 'Date' not in df_yesterday.columns:
            continue

        compare_cols = [col for col in df_today.columns if col != 'Date']

        today_keys = df_today[compare_cols].astype(str).apply(lambda row: '|'.join(row), axis=1)
        yesterday_keys = df_yesterday[compare_cols].astype(str).apply(lambda row: '|'.join(row), axis=1)

        only_yesterday_mask = ~yesterday_keys.isin(today_keys)
        only_yesterday_rows = df_yesterday.loc[only_yesterday_mask]

        if only_yesterday_rows.empty:
            continue

        sheet_data = pd.read_excel(output_path, sheet_name=sheet)
        for _, row in only_yesterday_rows.iterrows():
            mask = (sheet_data[compare_cols] == row[compare_cols].values).all(axis=1)
            match_indices = sheet_data[mask].index.tolist()
            for idx in match_indices:
                for col_idx in range(1, len(sheet_data.columns) + 1):
                    ws.cell(row=idx + 2, column=col_idx).fill = green_fill

    wb.save(output_path)

if __name__ == "__main__":
    convert_sarif_results("report/result/gitleaks.sarif", "gitleaks")
    convert_sarif_results("report/result/semgrep.sarif", "semgrep")
    convert_sarif_results("report/result/snyk-sast.sarif", "snyk-sast")
    convert_snyk_vuln_json("report/result/snyk.json", "snyk")
    convert_snyk_vuln_json("report/result/snyk-image.json", "snyk-image")
    convert_trivy_vuln("report/result/trivy-fs.json", "trivy-fs")
    convert_trivy_vuln("report/result/trivy-image.json", "trivy-image")
    convert_trivy_k8s("report/result/trivy-k8s.json", "trivy-k8s")
    merge_semgrep_and_snyk_to_sast()
    merge_trivy_and_snyk_to_sca()
    merge_trivy_and_snyk_to_image_scan()
    # merge_devsecops_reports('devsecops-report-02-04-25.xlsx', 'devsecops-report-03-04-25.xlsx', 'merged_output.xlsx')
